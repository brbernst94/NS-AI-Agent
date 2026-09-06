"""Document ingestion pipeline: PDF, Word, CSV, URL, and plain text.

Every ingest path accepts an optional `tags` dict which is merged into chunk
metadata. Recognised tags (system, module, doc_type, version, title, url,
tenant_id) are promoted to indexed columns by KnowledgeIndex.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# Chunk size in approximate characters (1 token ≈ 4 chars, ~500 tokens ≈ 2000 chars)
_CHUNK_SIZE = 2000
_CHUNK_OVERLAP = 200


def chunk_text(text: str, chunk_size: int = _CHUNK_SIZE, overlap: int = _CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks on paragraph boundaries where possible."""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []

    paragraphs = text.split("\n\n")
    chunks: list[str] = []
    current = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(current) + len(para) + 2 <= chunk_size:
            current = (current + "\n\n" + para).strip()
        else:
            if current:
                chunks.append(current)
            if len(para) > chunk_size:
                for i in range(0, len(para), chunk_size - overlap):
                    sub = para[i : i + chunk_size]
                    if sub.strip():
                        chunks.append(sub.strip())
                current = ""
            else:
                current = para

    if current:
        chunks.append(current)

    if overlap > 0 and len(chunks) > 1:
        overlapped: list[str] = [chunks[0]]
        for i in range(1, len(chunks)):
            tail = chunks[i - 1][-overlap:]
            overlapped.append((tail + " " + chunks[i]).strip())
        return overlapped

    return chunks


_chunk_text = chunk_text  # backwards-compatible alias


class DocumentIngester:
    """Ingest documents into the KnowledgeIndex from various formats."""

    def __init__(self, knowledge_index: Any, default_tags: dict[str, Any] | None = None) -> None:
        self.index = knowledge_index
        self.default_tags = dict(default_tags or {})

    def _meta(self, base: dict[str, Any], tags: dict[str, Any] | None) -> dict[str, Any]:
        meta = dict(self.default_tags)
        meta.update(base)
        if tags:
            meta.update({k: v for k, v in tags.items() if v is not None})
        return meta

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def ingest_pdf(self, file_path: str | Path, tags: dict[str, Any] | None = None) -> int:
        file_path = Path(file_path)
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(file_path))
        except Exception as exc:
            logger.error("Failed to read PDF '%s': %s", file_path, exc)
            return 0

        docs: list[dict[str, Any]] = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                continue
            for chunk_idx, chunk in enumerate(chunk_text(text)):
                docs.append({
                    "text": chunk,
                    "source": file_path.name,
                    "metadata": self._meta(
                        {"page": page_num, "chunk_index": chunk_idx, "file_type": "pdf",
                         "doc_type": "upload", "title": file_path.stem},
                        tags,
                    ),
                })

        added = self.index.add_documents(docs)
        logger.info("Ingested PDF '%s': %d chunks added.", file_path.name, added)
        return added

    # ------------------------------------------------------------------
    # Word (.docx)
    # ------------------------------------------------------------------

    def ingest_word(self, file_path: str | Path, tags: dict[str, Any] | None = None) -> int:
        file_path = Path(file_path)
        from docx import Document

        try:
            doc = Document(str(file_path))
        except Exception as exc:
            logger.error("Failed to read Word document '%s': %s", file_path, exc)
            return 0

        full_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())

        table_texts: list[str] = []
        for table in doc.tables:
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                table_texts.append("\n".join(rows))
        if table_texts:
            full_text += "\n\n" + "\n\n".join(table_texts)

        if not full_text.strip():
            logger.warning("No text extracted from Word document '%s'.", file_path.name)
            return 0

        docs = [
            {
                "text": chunk,
                "source": file_path.name,
                "metadata": self._meta(
                    {"chunk_index": i, "file_type": "docx", "doc_type": "upload",
                     "title": file_path.stem},
                    tags,
                ),
            }
            for i, chunk in enumerate(chunk_text(full_text))
        ]
        added = self.index.add_documents(docs)
        logger.info("Ingested Word doc '%s': %d chunks added.", file_path.name, added)
        return added

    # ------------------------------------------------------------------
    # CSV / Excel
    # ------------------------------------------------------------------

    def ingest_csv(self, file_path: str | Path, tags: dict[str, Any] | None = None) -> int:
        """Treat a CSV/Excel file as structured documentation, 20 rows per chunk."""
        file_path = Path(file_path)
        import pandas as pd

        try:
            if file_path.suffix.lower() in (".xlsx", ".xls"):
                df = pd.read_excel(str(file_path), dtype=str)
            else:
                df = pd.read_csv(str(file_path), dtype=str)
        except Exception as exc:
            logger.error("Failed to read CSV/Excel '%s': %s", file_path, exc)
            return 0

        df = df.fillna("")
        columns = list(df.columns)
        base = {"file_type": "csv", "doc_type": "upload", "title": file_path.stem}

        docs: list[dict[str, Any]] = [{
            "text": (
                f"File: {file_path.name}\n"
                f"Columns ({len(columns)}): {', '.join(columns)}\n"
                f"Total rows: {len(df)}"
            ),
            "source": file_path.name,
            "metadata": self._meta({**base, "chunk_type": "header"}, tags),
        }]

        batch_size = 20
        for batch_start in range(0, len(df), batch_size):
            batch = df.iloc[batch_start : batch_start + batch_size]
            lines = [f"Records {batch_start + 1}–{batch_start + len(batch)} from {file_path.name}:"]
            for _, row in batch.iterrows():
                parts = [f"{col}={val}" for col, val in row.items() if val]
                if parts:
                    lines.append("  { " + ", ".join(parts) + " }")
            docs.append({
                "text": "\n".join(lines),
                "source": file_path.name,
                "metadata": self._meta({**base, "chunk_type": "rows", "row_start": batch_start}, tags),
            })

        added = self.index.add_documents(docs)
        logger.info("Ingested CSV '%s': %d chunks added.", file_path.name, added)
        return added

    # ------------------------------------------------------------------
    # URL
    # ------------------------------------------------------------------

    def ingest_url(self, url: str, tags: dict[str, Any] | None = None) -> int:
        import requests
        from bs4 import BeautifulSoup

        try:
            resp = requests.get(url, timeout=30, headers={"User-Agent": "NS-AI-Agent/1.0"})
            resp.raise_for_status()
        except Exception as exc:
            logger.error("Failed to fetch URL '%s': %s", url, exc)
            return 0

        try:
            soup = BeautifulSoup(resp.content, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
                tag.decompose()
            main = soup.find("main") or soup.find("article") or soup.find("div", {"id": "content"})
            target = main if main else soup
            text = target.get_text(separator="\n", strip=True)
            title_tag = soup.find("title")
            title = title_tag.get_text(strip=True) if title_tag else None
        except Exception as exc:
            logger.error("Failed to parse HTML from '%s': %s", url, exc)
            return 0

        if not text.strip():
            logger.warning("No text extracted from URL '%s'.", url)
            return 0

        return self.ingest_text(
            text,
            source_name=url[:100],
            tags={"file_type": "url", "url": url, "title": title, "doc_type": "url", **(tags or {})},
        )

    # ------------------------------------------------------------------
    # Plain text
    # ------------------------------------------------------------------

    def ingest_text(self, text: str, source_name: str, tags: dict[str, Any] | None = None) -> int:
        if not text.strip():
            return 0
        docs = [
            {
                "text": chunk,
                "source": source_name,
                "metadata": self._meta({"chunk_index": i, "file_type": "text", "doc_type": "upload"}, tags),
            }
            for i, chunk in enumerate(chunk_text(text))
        ]
        added = self.index.add_documents(docs)
        logger.info("Ingested text '%s': %d chunks added.", source_name, added)
        return added

    # ------------------------------------------------------------------
    # File dispatcher
    # ------------------------------------------------------------------

    def ingest_file(
        self,
        file_path: str | Path,
        content: bytes | None = None,
        tags: dict[str, Any] | None = None,
    ) -> int:
        """Ingest a file by extension. If `content` is given, it is written to a temp file first."""
        file_path = Path(file_path)
        ext = file_path.suffix.lower()

        if content is not None:
            import tempfile

            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(content)
                tmp_path = Path(tmp.name)
            try:
                return self._dispatch(tmp_path, original_name=file_path.name, tags=tags)
            finally:
                tmp_path.unlink(missing_ok=True)
        return self._dispatch(file_path, original_name=file_path.name, tags=tags)

    def _dispatch(self, file_path: Path, original_name: str, tags: dict[str, Any] | None) -> int:
        ext = file_path.suffix.lower()
        tags = {"title": Path(original_name).stem, **(tags or {})}
        if ext == ".pdf":
            return self.ingest_pdf(file_path, tags)
        if ext in (".docx", ".doc"):
            return self.ingest_word(file_path, tags)
        if ext in (".csv", ".xlsx", ".xls"):
            return self.ingest_csv(file_path, tags)
        if ext in (".txt", ".md", ".rst"):
            text = file_path.read_text(encoding="utf-8", errors="replace")
            return self.ingest_text(text, original_name, tags)
        logger.warning("Unsupported file type: %s", ext)
        return 0
